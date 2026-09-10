"""
Bronze -> base -> gold pipeline for weekly 5-a-side scorecards.

Each source is one Google Sheets export (.xlsx) containing tabs:
  Template (skip), GW1..GWN (match data), plus derived summary tabs we
  don't trust and recompute instead (Captaincy, Analysis Tab, Apps, Awards).

Usage:
    python ingest.py <source_name> <path_to_xlsx> [<source_name> <path_to_xlsx> ...]
"""
import json
import re
import sys
from pathlib import Path

import duckdb
import openpyxl

ROOT = Path(__file__).parent
BRONZE = ROOT / "data" / "bronze"
GOLD = ROOT / "data" / "gold"
DB_PATH = ROOT / "data" / "league.duckdb"

GW_TAB_RE = re.compile(r"^GW(\d+)$", re.IGNORECASE)


def trim_trailing_none(row: tuple) -> list:
    row = list(row)
    while row and row[-1] is None:
        row.pop()
    return row


def load_workbook_tabs(xlsx_path: Path) -> dict[str, list[list]]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    tabs = {}
    for name in wb.sheetnames:
        ws = wb[name]
        tabs[name] = [trim_trailing_none(row) for row in ws.iter_rows(values_only=True)]
    return tabs


def dump_bronze(source: str, tabs: dict[str, list[list]]) -> None:
    out_dir = BRONZE / source
    out_dir.mkdir(parents=True, exist_ok=True)
    for tab_name, rows in tabs.items():
        safe_name = tab_name.replace("/", "_")
        with open(out_dir / f"{safe_name}.json", "w", encoding="utf-8") as f:
            json.dump(rows, f, default=str, indent=2)


def get(row: list, i: int, default=None):
    return row[i] if i < len(row) else default


def normalize_cs(value):
    # Source data inconsistency: some rows use "Y" (like the captain flag)
    # instead of a count in the clean-sheets column.
    if value == "Y":
        return 1
    return value


def clean_text(value):
    # Manual data entry leaves stray leading/trailing whitespace (e.g. "Jez "
    # vs "Jez"), which would otherwise fracture one player into two entries.
    return value.strip() if isinstance(value, str) else value


def parse_gw_tab(gw_number: int, rows: list[list], source: str) -> tuple[dict, list[dict]]:
    score_row = None
    for row in rows:
        if row and row[0] == "Score":
            score_row = row
            break
    if score_row is None:
        raise ValueError(f"{source} GW{gw_number}: no Score row found")

    match = {
        "source": source,
        "gw": gw_number,
        "team1_result": clean_text(get(score_row, 1)),
        "team1_goals": get(score_row, 3),
        "team2_result": clean_text(get(score_row, 6)),
        "team2_goals": get(score_row, 8),
        "venue": clean_text(get(score_row, 11)),
    }

    appearances = []
    for row in rows:
        if not row or row[0] == "Score" or row is rows[0]:
            continue
        # Team 1: name=col1, captain=col2, goals=col3, assists=col4, cs=col5
        # Data-entry slip recovery: col0 should only ever hold "Score" or the
        # GW header (both already skipped above); if col1 is blank but col0
        # has a name, the player's name was mistyped one cell to the left.
        team1_name = clean_text(get(row, 1))
        if team1_name is None:
            team1_name = clean_text(get(row, 0))
        if team1_name is not None:
            appearances.append({
                "source": source,
                "gw": gw_number,
                "team_side": 1,
                "player": team1_name,
                "is_captain": clean_text(get(row, 2)) == "Y",
                "goals": get(row, 3),
                "assists": get(row, 4),
                "clean_sheets": normalize_cs(get(row, 5)),
            })
        # Team 2: name=col6, captain=col7, goals=col8, assists=col9, cs=col10
        if clean_text(get(row, 6)) is not None:
            appearances.append({
                "source": source,
                "gw": gw_number,
                "team_side": 2,
                "player": clean_text(get(row, 6)),
                "is_captain": clean_text(get(row, 7)) == "Y",
                "goals": get(row, 8),
                "assists": get(row, 9),
                "clean_sheets": normalize_cs(get(row, 10)),
            })
    return match, appearances


def process_source(source: str, xlsx_path: Path, all_matches: list, all_appearances: list) -> None:
    tabs = load_workbook_tabs(xlsx_path)
    dump_bronze(source, tabs)

    gw_tabs = []
    for tab_name in tabs:
        m = GW_TAB_RE.match(tab_name.strip())
        if m:
            gw_tabs.append((int(m.group(1)), tab_name))
    gw_tabs.sort()

    for gw_number, tab_name in gw_tabs:
        match, appearances = parse_gw_tab(gw_number, tabs[tab_name], source)
        all_matches.append(match)
        all_appearances.extend(appearances)

        def player_view(a: dict) -> dict:
            return {k: v for k, v in a.items() if k not in ("source", "gw", "team_side")}

        gw_dir = GOLD / source
        gw_dir.mkdir(parents=True, exist_ok=True)
        record = {
            **match,
            "team1_players": [player_view(a) for a in appearances if a["team_side"] == 1],
            "team2_players": [player_view(a) for a in appearances if a["team_side"] == 2],
        }
        with open(gw_dir / f"GW{gw_number}.json", "w", encoding="utf-8") as f:
            json.dump(record, f, indent=2)

    source_appearances = sum(1 for a in all_appearances if a["source"] == source)
    print(f"{source}: parsed {len(gw_tabs)} gameweeks, {source_appearances} appearances")


def build_duckdb(all_matches: list[dict], all_appearances: list[dict]) -> None:
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    con = duckdb.connect(str(DB_PATH))
    import pandas as pd
    matches_df = pd.DataFrame(all_matches)
    appearances_df = pd.DataFrame(all_appearances)
    con.execute("CREATE OR REPLACE TABLE matches AS SELECT * FROM matches_df")
    con.execute("CREATE OR REPLACE TABLE appearances AS SELECT * FROM appearances_df")
    con.close()


def main():
    args = sys.argv[1:]
    if not args or len(args) % 2 != 0:
        print(__doc__)
        sys.exit(1)

    all_matches: list[dict] = []
    all_appearances: list[dict] = []

    for i in range(0, len(args), 2):
        source, path = args[i], Path(args[i + 1])
        process_source(source, path, all_matches, all_appearances)

    build_duckdb(all_matches, all_appearances)
    print(f"\nTotal: {len(all_matches)} matches, {len(all_appearances)} appearances across "
          f"{len(set(m['source'] for m in all_matches))} source(s)")
    print(f"DuckDB: {DB_PATH}")
    print(f"Gold JSON: {GOLD}")


if __name__ == "__main__":
    main()
